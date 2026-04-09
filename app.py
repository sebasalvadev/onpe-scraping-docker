"""
app.py — Consulta Miembro de Mesa ONPE
Lee DNIs del Excel, consulta el portal de ONPE con Selenium,
actualiza el Excel con los resultados y expone una interfaz web.
"""

import os
import time
import threading
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import Flask, jsonify
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, WebDriverException

app = Flask(__name__)

EXCEL_PATH = os.environ.get("EXCEL_PATH", "dnis.xlsx")
ONPE_URL   = "https://consultaelectoral.onpe.gob.pe/inicio"

# Estado global
_estado = {
    "running": False, "done": False,
    "current": 0, "total": 0,
    "resultados": [], "error": ""
}
_lock = threading.Lock()


def get_driver():
    opts = Options()
    opts.add_argument("--headless=new")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--disable-blink-features=AutomationControlled")
    opts.add_argument("--window-size=1280,900")
    opts.add_argument(
        "--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
    try:
        return webdriver.Chrome(options=opts)
    except Exception:
        pass
    for path in ["/usr/bin/chromedriver", "/usr/lib/chromium/chromedriver",
                 "/usr/lib/chromium-browser/chromedriver"]:
        if os.path.exists(path):
            return webdriver.Chrome(service=Service(path), options=opts)
    raise RuntimeError("No se encontró chromedriver")


def consultar_dni(driver, dni):
    resultado = {
        "dni": dni, "es_miembro": "NO",
        "nombre": "", "region": "", "provincia": "",
        "distrito": "", "direccion_local": "",
        "numero_mesa": "", "rol": "", "error": ""
    }
    try:
        driver.get(ONPE_URL)
        time.sleep(4)

        # Buscar campo DNI
        campo = None
        for by, sel in [
            (By.CSS_SELECTOR, "input[maxlength='8']"),
            (By.CSS_SELECTOR, "input[type='text']"),
            (By.CSS_SELECTOR, "input[type='number']"),
            (By.XPATH, "//input[contains(@placeholder,'DNI') or contains(@placeholder,'dni')]"),
            (By.XPATH, "//input"),
        ]:
            try:
                for el in driver.find_elements(by, sel):
                    if el.is_displayed() and el.is_enabled():
                        campo = el
                        break
                if campo:
                    break
            except Exception:
                continue

        if not campo:
            resultado["error"] = "No se encontró campo DNI"
            return resultado

        # Limpiar y escribir DNI de forma robusta
        driver.execute_script("arguments[0].value = '';", campo)
        campo.click()
        time.sleep(0.3)
        for char in dni:
            campo.send_keys(char)
            time.sleep(0.08)

        # Verificar que se escribió bien
        val = driver.execute_script("return arguments[0].value;", campo)
        if val != dni:
            driver.execute_script(f"arguments[0].value = '{dni}';", campo)
            driver.execute_script(
                "arguments[0].dispatchEvent(new Event('input', {{bubbles:true}}));"
                "arguments[0].dispatchEvent(new Event('change', {{bubbles:true}}));",
                campo
            )
        time.sleep(0.5)

        # Buscar botón consultar
        btn = None
        for by, sel in [
            (By.CSS_SELECTOR, "button[type='submit']"),
            (By.XPATH, "//button[contains(text(),'Consultar') or contains(text(),'CONSULTAR') or contains(text(),'Buscar')]"),
            (By.CSS_SELECTOR, "button.btn-primary, button.btn-danger, button.btn"),
            (By.XPATH, "//button"),
        ]:
            try:
                for el in driver.find_elements(by, sel):
                    if el.is_displayed() and el.is_enabled():
                        btn = el
                        break
                if btn:
                    break
            except Exception:
                continue

        if not btn:
            resultado["error"] = "No se encontró botón"
            return resultado

        try:
            btn.click()
        except Exception:
            driver.execute_script("arguments[0].click();", btn)

        # Esperar hasta 25s a que carguen los resultados
        PALABRAS_RESULTADO = [
            "eres presidente", "eres secretario", "eres tercer miembro",
            "eres suplente", "nombres y apellidos", "local de votación",
            "n° de mesa", "no eres miembro", "no figura", "no registra",
        ]
        for _ in range(50):
            time.sleep(0.5)
            try:
                body = driver.find_element(By.TAG_NAME, "body").text.lower()
                if any(p in body for p in PALABRAS_RESULTADO):
                    break
            except Exception:
                pass

        time.sleep(1)

        page_text = driver.find_element(By.TAG_NAME, "body").text
        page_lower = page_text.lower()

        # DEBUG
        with open("/app/debug_onpe.txt", "w", encoding="utf-8") as f:
            f.write(page_text)
        driver.save_screenshot("/app/debug_onpe.png")

        lineas = [l.strip() for l in page_text.split("\n") if l.strip()]

        # ── Detección de rol basada en el banner "ERES X" del portal ──
        if "eres presidente" in page_lower:
            resultado["rol"] = "Presidente de Mesa"
            resultado["es_miembro"] = "SÍ"
        elif "eres secretario" in page_lower:
            resultado["rol"] = "Secretario de Mesa"
            resultado["es_miembro"] = "SÍ"
        elif "eres tercer miembro" in page_lower:
            resultado["rol"] = "Tercer Miembro"
            resultado["es_miembro"] = "SÍ"
        elif "eres suplente" in page_lower:
            resultado["rol"] = "Suplente"
            resultado["es_miembro"] = "SÍ"
        else:
            resultado["rol"] = ""
            resultado["es_miembro"] = "NO"

        # ── Nombre: viene después de "Nombres y Apellidos" ──
        for i, linea in enumerate(lineas):
            if "nombres y apellidos" in linea.lower():
                if i + 1 < len(lineas):
                    resultado["nombre"] = lineas[i + 1]
                break

        # ── Región/Provincia/Distrito: formato "LIMA / LIMA / ATE" ──
        for i, linea in enumerate(lineas):
            if "región / provincia / distrito" in linea.lower() or "region / provincia / distrito" in linea.lower():
                if i + 1 < len(lineas):
                    partes = lineas[i + 1].split("/")
                    if len(partes) >= 3:
                        resultado["region"]    = partes[0].strip()
                        resultado["provincia"] = partes[1].strip()
                        resultado["distrito"]  = partes[2].strip()
                break

        # ── N° de Mesa ──
        for linea in lineas:
            if "n° de mesa:" in linea.lower():
                resultado["numero_mesa"] = linea.split(":")[-1].strip()
                break

        # ── Dirección del local ──
        for i, linea in enumerate(lineas):
            if "tu local de votación" in linea.lower():
                # Saltar la línea "ver Mapa" y tomar la siguiente con contenido real
                for j in range(i + 1, min(i + 4, len(lineas))):
                    if lineas[j].lower() not in ["ver mapa", "ver", "mapa"] and len(lineas[j]) > 5:
                        resultado["direccion_local"] = lineas[j]
                        break
                break

        if not resultado["nombre"] and not resultado["distrito"]:
            resultado["error"] = "Datos no encontrados: " + " | ".join(lineas[:4])

    except TimeoutException:
        resultado["error"] = "Timeout — portal no respondió"
    except Exception as e:
        resultado["error"] = str(e)[:120]
    return resultado


def leer_dnis():
    if not os.path.exists(EXCEL_PATH):
        return []
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
        dni_col = 2
        for col in ws.iter_cols(min_row=1, max_row=6, max_col=10):
            for cell in col:
                if cell.value and "dni" in str(cell.value).lower():
                    dni_col = cell.column
                    break
        dnis = []
        for row in ws.iter_rows(min_row=2, min_col=dni_col, max_col=dni_col):
            for cell in row:
                val = str(cell.value).strip() if cell.value else ""
                if val and val.isdigit():
                    dnis.append(val.zfill(8))
        return dnis
    except Exception as e:
        return []


def escribir_resultados(resultados):
    if not os.path.exists(EXCEL_PATH):
        return
    try:
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
        dni_col = 2
        for col in ws.iter_cols(min_row=1, max_row=6, max_col=10):
            for cell in col:
                if cell.value and "dni" in str(cell.value).lower():
                    dni_col = cell.column
                    break
        start_col = dni_col + 1
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        headers = ["¿Miembro?", "Nombre", "Región", "Provincia",
                   "Distrito", "Dirección del Local", "N° Mesa", "Rol", "Nota"]
        for i, h in enumerate(headers):
            c = ws.cell(row=4, column=start_col + i, value=h)
            c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            c.fill = PatternFill("solid", start_color="C0392B")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
        ws.row_dimensions[4].height = 28
        bg_alt = ["FFFFFF", "F2F3F4"]
        for i, res in enumerate(resultados):
            row = 5 + i
            bg = bg_alt[i % 2]
            color_m = "1E8449" if res["es_miembro"] == "SÍ" else "C0392B"
            def wd(offset, val, color="000000", bold=False, center=False, r=row, b=bg):
                c = ws.cell(row=r, column=start_col + offset, value=val)
                c.font = Font(name="Arial", size=10, color=color, bold=bold)
                c.fill = PatternFill("solid", start_color=b)
                c.alignment = Alignment(horizontal="center" if center else "left", vertical="center")
                c.border = border
            wd(0, res["es_miembro"], color_m, bold=True, center=True)
            wd(1, res["nombre"])
            wd(2, res["region"])
            wd(3, res["provincia"])
            wd(4, res["distrito"])
            wd(5, res["direccion_local"])
            wd(6, res["numero_mesa"], center=True)
            wd(7, res["rol"])
            wd(8, res["error"], "E74C3C" if res["error"] else "27AE60")
            ws.row_dimensions[row].height = 20
        anchos = [12, 28, 14, 14, 16, 42, 10, 20, 30]
        for i, w in enumerate(anchos):
            ws.column_dimensions[openpyxl.utils.get_column_letter(start_col + i)].width = w
        wb.save(EXCEL_PATH)
    except Exception as e:
        print(f"Error guardando Excel: {e}")


def procesar():
    dnis = leer_dnis()
    if not dnis:
        with _lock:
            _estado["running"] = False
            _estado["error"] = (
                f"No se encontraron DNIs en {EXCEL_PATH}. "
                "Verifica que el archivo existe y tiene DNIs en la columna B."
            )
        return
    with _lock:
        _estado["total"] = len(dnis)
    try:
        driver = get_driver()
    except Exception as e:
        with _lock:
            _estado["running"] = False
            _estado["error"] = f"No se pudo iniciar el navegador: {e}"
        return
    resultados = []
    for i, dni in enumerate(dnis):
        with _lock:
            _estado["current"] = i + 1
        
        # Intentar hasta 3 veces si falla
        res = consultar_dni(driver, dni)
        if not res["nombre"] and not res["distrito"]:
            time.sleep(5)
            res = consultar_dni(driver, dni)
        if not res["nombre"] and not res["distrito"]:
            time.sleep(8)
            res = consultar_dni(driver, dni)
        
        resultados.append(res)
        time.sleep(3)  # Pausa entre DNIs para no saturar el portal
    driver.quit()
    escribir_resultados(resultados)
    with _lock:
        _estado["running"] = False
        _estado["done"] = True
        _estado["resultados"] = resultados


HTML = """<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Consulta Miembro de Mesa — ONPE</title>
<style>
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:Arial,sans-serif;background:#f0f2f5;min-height:100vh}
header{background:#c0392b;color:#fff;padding:18px 32px}
header h1{font-size:20px}
header p{font-size:12px;opacity:.85;margin-top:3px}
.container{max-width:1100px;margin:28px auto;padding:0 16px}
.card{background:#fff;border-radius:10px;box-shadow:0 2px 12px rgba(0,0,0,.08);padding:28px;margin-bottom:24px}
.card h2{font-size:16px;color:#2c3e50;margin-bottom:16px;padding-bottom:8px;border-bottom:2px solid #e8e8e8}
.info-box{background:#eaf4fb;border-left:4px solid #2980b9;padding:12px 16px;border-radius:4px;font-size:13px;color:#1a5276;margin-bottom:16px;line-height:1.6}
.btn{padding:10px 24px;border-radius:6px;border:none;cursor:pointer;font-size:14px;font-weight:bold;transition:.2s}
.btn-primary{background:#c0392b;color:#fff}.btn-primary:hover{background:#a93226}
.btn-primary:disabled{background:#bbb;cursor:not-allowed}
.status{margin-top:14px;padding:12px 16px;border-radius:6px;font-size:13px;display:none;line-height:1.5}
.running{background:#fef9e7;border:1px solid #f39c12;color:#7d6608}
.done{background:#eafaf1;border:1px solid #27ae60;color:#1e8449}
.err{background:#fdf2f2;border:1px solid #e74c3c;color:#922b21}
.tbl-wrap{overflow-x:auto;margin-top:16px}
table{width:100%;border-collapse:collapse;font-size:13px}
th{background:#c0392b;color:#fff;padding:10px 12px;text-align:left}
td{padding:9px 12px;border-bottom:1px solid #eee;vertical-align:middle}
tr:nth-child(even){background:#f9f9f9}
tr:hover{background:#fef5f5}
.si{color:#1e8449;font-weight:bold}.no{color:#c0392b;font-weight:bold}
.empty{text-align:center;color:#aaa;padding:32px;font-size:15px}
footer{text-align:center;color:#aaa;font-size:11px;padding:20px}
</style>
</head>
<body>
<header>
  <h1>🗳️ Consulta Miembro de Mesa — ONPE</h1>
  <p>Laboratorio 04 · Contenedores Docker · Desarrollo de Soluciones en la Nube</p>
</header>
<div class="container">
  <div class="card">
    <h2>⚙️ Ejecutar Consulta</h2>
    <div class="info-box">
      El sistema leerá los DNIs del archivo <strong>dnis.xlsx</strong>,
      consultará cada uno en el portal oficial de la ONPE
      (<a href="https://consultaelectoral.onpe.gob.pe/inicio" target="_blank">consultaelectoral.onpe.gob.pe</a>)
      y actualizará el Excel con los resultados automáticamente.<br><br>
      ⚠️ <strong>Cada DNI tarda ~15-30 segundos.</strong> Para 5 DNIs espera entre 2 y 3 minutos.
    </div>
    <button class="btn btn-primary" id="btnConsultar" onclick="iniciar()">
      🔍 Consultar todos los DNIs
    </button>
    <div class="status" id="statusBox"></div>
  </div>
  <div class="card">
    <h2>📋 Resultados</h2>
    <div id="resultado"><p class="empty">Presiona el botón para iniciar la consulta.</p></div>
  </div>
</div>
<footer>Caso 2 — ONPE Electoral Consultation System · Docker Lab 04</footer>
<script>
let poll=null;
function iniciar(){
  const btn=document.getElementById('btnConsultar');
  btn.disabled=true;
  mostrarStatus('running','⏳ Iniciando consulta... espera un momento.');
  fetch('/consultar',{method:'POST'}).then(r=>r.json()).then(d=>{
    if(d.started){poll=setInterval(verificar,4000);}
    else{mostrarStatus('err','❌ '+d.msg);btn.disabled=false;}
  });
}
function verificar(){
  fetch('/status').then(r=>r.json()).then(d=>{
    const btn=document.getElementById('btnConsultar');
    if(d.error){clearInterval(poll);mostrarStatus('err','❌ '+d.error);btn.disabled=false;}
    else if(d.running){mostrarStatus('running','⏳ Consultando DNI '+d.current+' de '+d.total+'... espera.');}
    else if(d.done){clearInterval(poll);mostrarStatus('done','✅ Completado — '+d.total+' DNIs procesados. Excel actualizado.');btn.disabled=false;renderTabla(d.resultados);}
  });
}
function mostrarStatus(tipo,msg){
  const box=document.getElementById('statusBox');
  box.style.display='block';box.className='status '+tipo;box.innerHTML=msg;
}
function renderTabla(data){
  if(!data||!data.length)return;
  let html='<div class="tbl-wrap"><table><thead><tr><th>#</th><th>DNI</th><th>¿Miembro?</th><th>Nombre</th><th>Región</th><th>Provincia</th><th>Distrito</th><th>Dirección del Local</th><th>N° Mesa</th><th>Rol</th><th>Nota</th></tr></thead><tbody>';
  data.forEach((r,i)=>{
    const cls=r.es_miembro==='SÍ'?'si':'no';
    const ico=r.es_miembro==='SÍ'?'✅ SÍ':'❌ NO';
    html+=`<tr><td>${i+1}</td><td><strong>${r.dni}</strong></td><td class="${cls}">${ico}</td><td>${r.nombre||'—'}</td><td>${r.region||'—'}</td><td>${r.provincia||'—'}</td><td>${r.distrito||'—'}</td><td>${r.direccion_local||'—'}</td><td>${r.numero_mesa||'—'}</td><td>${r.rol||'—'}</td><td style="font-size:11px;color:#888">${r.error||'OK'}</td></tr>`;
  });
  html+='</tbody></table></div>';
  document.getElementById('resultado').innerHTML=html;
}
fetch('/status').then(r=>r.json()).then(d=>{if(d.done&&d.resultados&&d.resultados.length)renderTabla(d.resultados);});
</script>
</body>
</html>"""


@app.route('/')
def index():
    return HTML

@app.route('/consultar', methods=['POST'])
def consultar():
    with _lock:
        if _estado["running"]:
            return jsonify({"started": False, "msg": "Ya hay una consulta en curso"})
        _estado.update({"running": True, "done": False, "current": 0,
                        "total": 0, "resultados": [], "error": ""})
    threading.Thread(target=procesar, daemon=True).start()
    return jsonify({"started": True})

@app.route('/status')
def status():
    with _lock:
        return jsonify(dict(_estado))

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=False)